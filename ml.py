import os.path
import base64
import re
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from email import message_from_bytes
from docx import Document

# If modifying scopes, delete token.json
SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']

def authenticate_gmail():
    creds = None
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file('credentials.json', SCOPES)
            creds = flow.run_local_server(port=0)
        with open('token.json', 'w') as token:
            token.write(creds.to_json())
    return build('gmail', 'v1', credentials=creds)

def fetch_emails(service, senders, max_results=100):
    all_messages = []
    for sender in senders:
        query = f'from:{sender}'
        response = service.users().messages().list(userId='me', q=query, maxResults=max_results).execute()
        messages = response.get('messages', [])
        for msg in messages:
            msg_detail = service.users().messages().get(userId='me', id=msg['id'], format='raw').execute()
            raw_msg = base64.urlsafe_b64decode(msg_detail['raw'].encode('ASCII'))
            mime_msg = message_from_bytes(raw_msg)
            subject = mime_msg['subject']
            date = mime_msg['date']
            body = extract_body(mime_msg)
            all_messages.append((subject, date, body))
    return all_messages

def extract_body(mime_msg):
    if mime_msg.is_multipart():
        for part in mime_msg.walk():
            content_type = part.get_content_type()
            if content_type == 'text/plain':
                return part.get_payload(decode=True).decode(errors='ignore')
    else:
        return mime_msg.get_payload(decode=True).decode(errors='ignore')
    return ""

def create_combined_doc(messages, output_file='combined_emails.docx'):
    doc = Document()
    doc.add_heading('Combined Emails from Specific Senders', 0)
    for subject, date, body in messages:
        doc.add_heading(subject or '(No Subject)', level=2)
        doc.add_paragraph(f"Date: {date}")
        doc.add_paragraph(body)
        doc.add_paragraph("\n---\n")
    doc.save(output_file)

if __name__ == '__main__':
    specific_senders = ['example1@gmail.com', 'example2@gmail.com']  # Replace with real addresses
    service = authenticate_gmail()
    emails = fetch_emails(service, specific_senders)
    create_combined_doc(emails)
    print("✅ Emails combined and saved to 'combined_emails.docx'")
